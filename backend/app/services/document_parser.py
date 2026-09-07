import os
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

def _clean_text(text: str) -> str:
    """Normalize extracted text by safely stripping redundant spaces and empty lines."""
    if not text:
        return ""
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    return "\n".join(lines)

def extract_pdf_text(filepath: str) -> str:
    text = ""
    try:
        doc = fitz.open(filepath)
        for page in doc:
            text += page.get_text() + "\n"
        doc.close()
    except Exception as e:
        raise ValueError(f"Failed to read PDF: {e}")
    return text

def extract_docx_text(filepath: str) -> str:
    text = []
    try:
        doc = Document(filepath)
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text.strip())
        
        # Also extract from basic tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text.strip())
    except Exception as e:
        raise ValueError(f"Failed to read DOCX: {e}")
    return "\n".join(text)

def extract_pptx_text(filepath: str) -> str:
    text = []
    try:
        prs = Presentation(filepath)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text.append(shape.text.strip())
    except Exception as e:
        raise ValueError(f"Failed to read PPTX: {e}")
    return "\n".join(text)

def extract_txt_text(filepath: str) -> str:
    try:
        with open(filepath, 'rb') as f:
            raw = f.read()
            # Standard utf-8 decode, replace offending characters
            return raw.decode('utf-8', errors='replace')
    except Exception as e:
        raise ValueError(f"Failed to read TXT: {e}")

async def extract_job_description(file_path: str, filename: str) -> dict:
    """
    Given a local file path, determines type and extracts normalized text.
    Returns dictionary with text payload and metadata.
    """
    ext = os.path.splitext(filename.lower())[1]
    
    if ext == '.pdf':
        text = extract_pdf_text(file_path)
        mime = "application/pdf"
    elif ext == '.docx':
        text = extract_docx_text(file_path)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == '.pptx':
        text = extract_pptx_text(file_path)
        mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif ext == '.txt':
        text = extract_txt_text(file_path)
        mime = "text/plain"
    else:
        raise ValueError(f"Unsupported file extension: {ext}. Only PDF, DOCX, PPTX, and TXT are supported.")
        
    cleaned_text = _clean_text(text)
    
    if not cleaned_text:
        if ext == '.pdf':
            raise ValueError("No selectable text was found in this file. Please paste the Job Description manually or upload a text-based PDF.")
        else:
            raise ValueError("The file appears to be empty or contains no readable text.")

    return {
        "filename": filename,
        "file_type": mime,
        "extracted_text": cleaned_text,
        "character_count": len(cleaned_text)
    }
